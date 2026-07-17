"""Document extraction and paragraph-aware chunking for Tracefold."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown"}


@dataclass(frozen=True)
class TextBlock:
    text: str
    page: int | None = None
    section: str | None = None


@dataclass(frozen=True)
class ParsedChunk:
    text: str
    ordinal: int
    page_start: int | None = None
    page_end: int | None = None
    section: str | None = None
    token_count: int = 0


class DocumentParser:
    """Extract PDF, DOCX, TXT, and Markdown while retaining source locations."""

    def parse(
        self,
        file_path: str | Path,
        *,
        chunk_size: int = 240,
        overlap: int = 45,
    ) -> list[ParsedChunk]:
        path = Path(file_path)
        extension = path.suffix.lower()
        if extension not in SUPPORTED_EXTENSIONS:
            supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
            raise ValueError(f"Unsupported file type '{extension or 'unknown'}'. Supported: {supported}.")

        if extension == ".pdf":
            blocks = self._pdf_blocks(path)
        elif extension == ".docx":
            blocks = self._docx_blocks(path)
        elif extension in {".md", ".markdown"}:
            blocks = self._markdown_blocks(path)
        else:
            blocks = self._text_blocks(path)

        if not any(block.text.strip() for block in blocks):
            if extension == ".pdf":
                raise ValueError(
                    "No selectable text was found. This may be a scanned PDF; OCR is not supported."
                )
            raise ValueError("No readable text was found in this document.")
        return self.chunk_blocks(blocks, chunk_size=chunk_size, overlap=overlap)

    def chunk_blocks(
        self,
        blocks: list[TextBlock],
        *,
        chunk_size: int = 240,
        overlap: int = 45,
    ) -> list[ParsedChunk]:
        if chunk_size <= 0 or overlap < 0 or overlap >= chunk_size:
            raise ValueError("chunk_size must be positive and greater than overlap.")

        units: list[tuple[str, int | None, str | None]] = []
        for block in blocks:
            for paragraph in self._paragraphs(block.text):
                units.append((paragraph, block.page, block.section))

        chunks: list[ParsedChunk] = []
        cursor = 0
        while cursor < len(units):
            selected: list[tuple[str, int | None, str | None]] = []
            token_total = 0
            index = cursor
            section_boundary = False
            while index < len(units):
                paragraph, page, section = units[index]
                if selected and section != selected[0][2]:
                    section_boundary = True
                    break
                tokens = self._tokens(paragraph)
                if selected and token_total + len(tokens) > chunk_size:
                    break
                if not selected and len(tokens) > chunk_size:
                    # Preserve normal paragraphs, but split pathological blocks at token boundaries.
                    selected.append((" ".join(tokens[:chunk_size]), page, section))
                    remainder = " ".join(tokens[max(0, chunk_size - overlap) :])
                    units[index] = (remainder, page, section)
                    token_total = chunk_size
                    break
                selected.append((paragraph, page, section))
                token_total += len(tokens)
                index += 1

            pages = [item[1] for item in selected if item[1] is not None]
            sections = [item[2] for item in selected if item[2]]
            chunks.append(
                ParsedChunk(
                    text="\n\n".join(item[0] for item in selected),
                    ordinal=len(chunks),
                    page_start=min(pages) if pages else None,
                    page_end=max(pages) if pages else None,
                    section=sections[0] if sections else None,
                    token_count=token_total,
                )
            )

            if index >= len(units):
                break
            if section_boundary:
                cursor = index
                continue
            # Walk backward over whole paragraphs until roughly `overlap` tokens are retained.
            retained = 0
            next_cursor = index
            while next_cursor > cursor and retained < overlap:
                next_cursor -= 1
                retained += len(self._tokens(units[next_cursor][0]))
            cursor = next_cursor if next_cursor > cursor else index
        return chunks

    def _pdf_blocks(self, path: Path) -> list[TextBlock]:
        import fitz

        document = fitz.open(path)
        try:
            return [
                TextBlock(text=page.get_text("text").strip(), page=page_number + 1)
                for page_number, page in enumerate(document)
                if page.get_text("text").strip()
            ]
        finally:
            document.close()

    def _docx_blocks(self, path: Path) -> list[TextBlock]:
        from docx import Document

        document = Document(path)
        blocks: list[TextBlock] = []
        current_section: str | None = None
        for paragraph in document.paragraphs:
            text = self._clean(paragraph.text)
            if not text:
                continue
            if paragraph.style and paragraph.style.name.lower().startswith("heading"):
                current_section = text
            else:
                blocks.append(TextBlock(text=text, section=current_section))
        for table_number, table in enumerate(document.tables, start=1):
            rows = []
            for row in table.rows:
                cells = [self._clean(cell.text) for cell in row.cells]
                rows.append(" | ".join(cells))
            if rows:
                blocks.append(TextBlock(text="\n".join(rows), section=f"Table {table_number}"))
        return blocks

    def _markdown_blocks(self, path: Path) -> list[TextBlock]:
        blocks: list[TextBlock] = []
        current_section: str | None = None
        buffer: list[str] = []
        for line in self._read_text(path).splitlines():
            heading = re.match(r"^#{1,6}\s+(.+?)\s*#*$", line.strip())
            if heading:
                if buffer:
                    blocks.append(TextBlock(text="\n".join(buffer), section=current_section))
                    buffer = []
                current_section = self._clean(heading.group(1))
            else:
                buffer.append(line)
        if buffer:
            blocks.append(TextBlock(text="\n".join(buffer), section=current_section))
        return blocks

    def _text_blocks(self, path: Path) -> list[TextBlock]:
        return [TextBlock(text=self._read_text(path))]

    def _read_text(self, path: Path) -> str:
        for encoding in ("utf-8-sig", "utf-8", "cp1254", "latin-1"):
            try:
                return path.read_text(encoding=encoding)
            except UnicodeDecodeError:
                continue
        raise ValueError("The text document could not be decoded.")

    @staticmethod
    def _paragraphs(text: str) -> list[str]:
        normalized = text.replace("\r\n", "\n").replace("\r", "\n")
        paragraphs = re.split(r"\n\s*\n|(?<=\.)\s*\n", normalized)
        return [DocumentParser._clean(item) for item in paragraphs if item.strip()]

    @staticmethod
    def _tokens(text: str) -> list[str]:
        # A deterministic approximation keeps chunking independent from a cloud tokenizer.
        return re.findall(r"\w+(?:[-'][\w]+)*|[^\w\s]", text, flags=re.UNICODE)

    @staticmethod
    def _clean(value: str) -> str:
        return re.sub(r"[ \t]+", " ", value).strip()
